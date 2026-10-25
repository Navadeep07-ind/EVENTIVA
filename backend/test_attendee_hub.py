import requests
import uuid
import io
import openpyxl
from datetime import datetime, timezone, timedelta

from app.database import SessionLocal
from app import models
from app.core.security import get_password_hash, create_access_token

BASE_URL = "http://127.0.0.1:8000"

def setup_test_environment():
    db = SessionLocal()
    try:
        # Create Organizer A
        user_a = db.query(models.User).filter(models.User.email == "hub_organizer_a@eventiva.test").first()
        if not user_a:
            user_a = models.User(
                email="hub_organizer_a@eventiva.test",
                full_name="Hub Organizer Alpha",
                hashed_password=get_password_hash("Password123!"),
                is_verified=True
            )
            db.add(user_a)
        else:
            user_a.is_verified = True

        # Create Organizer B
        user_b = db.query(models.User).filter(models.User.email == "hub_organizer_b@eventiva.test").first()
        if not user_b:
            user_b = models.User(
                email="hub_organizer_b@eventiva.test",
                full_name="Hub Organizer Beta",
                hashed_password=get_password_hash("Password123!"),
                is_verified=True
            )
            db.add(user_b)
        else:
            user_b.is_verified = True

        db.commit()
        db.refresh(user_a)
        db.refresh(user_b)

        token_a = create_access_token({"sub": user_a.email})
        token_b = create_access_token({"sub": user_b.email})
        headers_a = {"Authorization": f"Bearer {token_a}"}
        headers_b = {"Authorization": f"Bearer {token_b}"}

        # Event 1 (Owned by Organizer A) with custom form_schema
        event_1 = models.Event(
            title="Noventrix Tech Conclave 2026",
            description="Flagship Technical Symposium",
            category="Technology",
            venue="Auditorium Hall A",
            status="Published",
            organizer_id=user_a.id,
            form_schema=[
                {"id": "track_choice", "label": "Technical Track", "type": "dropdown"},
                {"id": "github_url", "label": "GitHub Profile", "type": "url"}
            ]
        )

        # Event 2 (Owned by Organizer B)
        event_2 = models.Event(
            title="FinTech Global Summit 2026",
            description="Finance Leadership",
            category="Business",
            venue="Convention Hall B",
            status="Published",
            organizer_id=user_b.id
        )

        db.add_all([event_1, event_2])
        db.commit()
        db.refresh(event_1)
        db.refresh(event_2)

        return user_a, user_b, event_1.id, event_2.id, headers_a, headers_b
    finally:
        db.close()

def run_tests():
    print("==================================================")
    print("EVENTIVA ATTENDEE OPERATIONS HUB & CHECK-IN SUITE")
    print("==================================================")

    _, _, event1_id, event2_id, headers_a, headers_b = setup_test_environment()

    # Pre-populate attendees for Event 1 (including potential formula injection payloads)
    attendee_1_payload = {
        "full_name": "Rohan Verma",
        "email": f"rohan_{uuid.uuid4().hex[:6]}@example.com",
        "college": "TKR College of Engineering and Technology",
        "custom_answers": {
            "track_choice": "Artificial Intelligence",
            "github_url": "https://github.com/rohanv",
            "phone": "+91 9123456780"
        }
    }
    attendee_2_payload = {
        "full_name": "Priya Patel",
        "email": f"priya_{uuid.uuid4().hex[:6]}@example.com",
        "college": "VNR VJIET",
        "custom_answers": {
            "track_choice": "Cybersecurity",
            "github_url": "https://github.com/priyap"
        }
    }
    # Formula injection test payload
    attendee_injection_payload = {
        "full_name": "=1+1 Formula Test",
        "email": f"injection_{uuid.uuid4().hex[:6]}@example.com",
        "college": "@SUM(A1:A10)",
        "custom_answers": {
            "track_choice": "+100 Bonus Track",
            "github_url": "-cmd|' /C calc'!A0"
        }
    }

    res1 = requests.post(f"{BASE_URL}/api/events/{event1_id}/register", json=attendee_1_payload)
    assert res1.status_code == 200, f"Registration 1 failed: {res1.text}"
    ticket_1 = res1.json()["ticket_code"]
    reg1_id = res1.json()["id"]

    res2 = requests.post(f"{BASE_URL}/api/events/{event1_id}/register", json=attendee_2_payload)
    assert res2.status_code == 200, f"Registration 2 failed: {res2.text}"
    ticket_2 = res2.json()["ticket_code"]
    reg2_id = res2.json()["id"]

    res_inj = requests.post(f"{BASE_URL}/api/events/{event1_id}/register", json=attendee_injection_payload)
    assert res_inj.status_code == 200, f"Registration injection failed: {res_inj.text}"

    # Pre-populate attendee for Event 2 (Owned by Organizer B)
    attendee_evt2_payload = {
        "full_name": "Ananya Reddy",
        "email": f"ananya_{uuid.uuid4().hex[:6]}@example.com",
        "college": "Osmania University"
    }
    res_evt2 = requests.post(f"{BASE_URL}/api/events/{event2_id}/register", json=attendee_evt2_payload)
    assert res_evt2.status_code == 200
    ticket_evt2 = res_evt2.json()["ticket_code"]

    print("[SETUP PASS] Test attendees registered on Event 1 and Event 2.")

    # 1. Authorized organizer can list event registrations
    res = requests.get(f"{BASE_URL}/api/events/{event1_id}/registrations", headers=headers_a)
    assert res.status_code == 200, f"List registrations failed: {res.text}"
    attendees_a = res.json()
    assert len(attendees_a) >= 3, "Organizer A must retrieve registered attendees"
    print(f"[TEST 1 PASS] Authorized Organizer A retrieved {len(attendees_a)} registrations for Event 1.")

    # 2. Unauthorized user cannot access attendee information
    res = requests.get(f"{BASE_URL}/api/events/{event1_id}/registrations", headers=headers_b)
    assert res.status_code == 403, f"Expected 403 Forbidden for unauthorized user, got {res.status_code}"
    print("[TEST 2 PASS] Security Enforced: Organizer B received 403 Forbidden when accessing Event 1 registrations.")

    # 3. Organizer cannot modify registrations belonging to another event
    res = requests.put(
        f"{BASE_URL}/api/events/{event1_id}/registrations/{reg1_id}",
        json={"status": "Approved"},
        headers=headers_b # Organizer B trying to edit Organizer A's attendee
    )
    assert res.status_code == 403, f"Expected 403 Forbidden, got {res.status_code}"
    print("[TEST 3 PASS] Security Enforced: Organizer B received 403 Forbidden when attempting to modify Event 1 registration.")

    # Status update by authorized owner
    res_status = requests.put(
        f"{BASE_URL}/api/events/{event1_id}/registrations/{reg1_id}",
        json={"status": "Approved"},
        headers=headers_a
    )
    assert res_status.status_code == 200
    assert res_status.json()["status"] == "Approved"
    print(f"[STATUS PASS] Organizer A updated registration {reg1_id} status to 'Approved'.")

    # 4. Valid ticket can be checked in exactly once
    checkin_payload = {"ticket_code": ticket_1}
    res = requests.post(f"{BASE_URL}/api/events/{event1_id}/check-in", json=checkin_payload, headers=headers_a)
    assert res.status_code == 200, f"Check-in failed: {res.text}"
    checkin_data = res.json()
    assert checkin_data["status"] == "success", "Check in status must be success"
    assert checkin_data["checked_in"] is True
    assert checkin_data["is_duplicate"] is False
    assert checkin_data["attendee_name"] == "Rohan Verma"
    print(f"[TEST 4 PASS] Valid ticket {ticket_1} successfully admitted at the venue.")

    # 5. Repeated check-in returns an already-checked-in response
    res_repeat = requests.post(f"{BASE_URL}/api/events/{event1_id}/check-in", json=checkin_payload, headers=headers_a)
    assert res_repeat.status_code == 200
    repeat_data = res_repeat.json()
    assert repeat_data["status"] == "already_checked_in", "Must report already checked in"
    assert repeat_data["is_duplicate"] is True
    print(f"[TEST 5 PASS] Duplicate Prevention: Repeated scan reported already_checked_in for {ticket_1}.")

    # 6. Cancelled/rejected/revoked ticket cannot be checked in
    # Organizer A cancels Registration 2
    res_cancel = requests.put(
        f"{BASE_URL}/api/events/{event1_id}/registrations/{reg2_id}",
        json={"status": "Cancelled"},
        headers=headers_a
    )
    assert res_cancel.status_code == 200
    assert res_cancel.json()["status"] == "Cancelled"

    # Attempt to check in cancelled ticket
    res_check_cancelled = requests.post(
        f"{BASE_URL}/api/events/{event1_id}/check-in",
        json={"ticket_code": ticket_2},
        headers=headers_a
    )
    assert res_check_cancelled.status_code == 400, "Cancelled ticket check-in must be rejected with 400"
    assert "denied" in res_check_cancelled.text.lower() or "cancelled" in res_check_cancelled.text.lower()
    print(f"[TEST 6 PASS] Revocation Enforced: Cancelled ticket {ticket_2} was denied entry with HTTP 400.")

    # 7. Ticket belonging to another event is rejected
    res_wrong_event = requests.post(
        f"{BASE_URL}/api/events/{event1_id}/check-in",
        json={"ticket_code": ticket_evt2}, # Belongs to Event 2!
        headers=headers_a
    )
    assert res_wrong_event.status_code == 400, "Ticket for another event must be rejected with 400"
    assert "different event" in res_wrong_event.text.lower()
    print(f"[TEST 7 PASS] Event Isolation: Ticket from Event 2 rejected when scanned at Event 1 entrance.")

    # 8. Public ticket access does not perform check-in
    # Check attendee 2's ticket pass
    res_pub = requests.get(f"{BASE_URL}/api/tickets/{ticket_2}")
    assert res_pub.status_code == 200
    pub_data = res_pub.json()
    assert pub_data["checked_in"] is False, "Public view must never check attendee in"
    # Ensure database record remains checked_in = False
    db = SessionLocal()
    reg2_db = db.query(models.Registration).filter(models.Registration.ticket_code == ticket_2).first()
    assert reg2_db.checked_in is False, "Database record must remain unchecked"
    db.close()
    print(f"[TEST 8 PASS] Public ticket access is strictly read-only and does NOT perform check-in.")

    # 9. Excel export contains actual database registrations
    res_export = requests.get(f"{BASE_URL}/api/events/{event1_id}/registrations/export", headers=headers_a)
    assert res_export.status_code == 200, f"Export failed: {res_export.status_code}"
    assert "spreadsheetml" in res_export.headers.get("content-type", "")
    assert res_export.headers.get("content-disposition", "").endswith('.xlsx"')

    wb = openpyxl.load_workbook(io.BytesIO(res_export.content))
    ws = wb.active
    assert ws.title == "Participants"
    header_row = [cell.value for cell in ws[1]]
    assert "Ticket Code" in header_row
    assert "Full Name" in header_row
    assert "Email" in header_row
    assert "Checked In" in header_row

    # Collect all attendee names from exported sheet
    exported_names = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "Rohan Verma" in exported_names, "Export must contain registered attendee Rohan Verma"
    assert "Priya Patel" in exported_names, "Export must contain registered attendee Priya Patel"
    print(f"[TEST 9 PASS] Excel Export generated valid binary .xlsx workbook with {len(exported_names)} attendee rows.")

    # 10. Excel export includes dynamic custom answers
    assert "Technical Track" in header_row, "Export must include custom question 'Technical Track'"
    assert "GitHub Profile" in header_row, "Export must include custom question 'GitHub Profile'"
    track_col_idx = header_row.index("Technical Track")
    
    # Verify Rohan's track is 'Artificial Intelligence'
    rohan_row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == "Rohan Verma")
    assert rohan_row[track_col_idx] == "Artificial Intelligence"
    print(f"[TEST 10 PASS] Dynamic custom form schema columns ('Technical Track', 'GitHub Profile') present with correct attendee responses.")

    # 11. Spreadsheet formula injection is prevented
    # Look for the row with injection payload
    inj_row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if "Formula Test" in str(r[1]))
    # Cell values beginning with =, +, -, @ must be prepended with a single quote '
    name_val = str(inj_row[1])
    college_val = str(inj_row[3])
    assert name_val.startswith("'="), f"Formula injection not escaped in name: {name_val}"
    assert college_val.startswith("'@"), f"Formula injection not escaped in college: {college_val}"
    print(f"[TEST 11 PASS] Security Verified: Formula injection payloads safely escaped with prepended quote ('=, '@).")

    # 12. Search and Filter endpoint verification
    res_search = requests.get(f"{BASE_URL}/api/events/{event1_id}/registrations?search=Rohan", headers=headers_a)
    assert res_search.status_code == 200
    search_results = res_search.json()
    assert len(search_results) == 1 and search_results[0]["full_name"] == "Rohan Verma"

    res_filter_checked = requests.get(f"{BASE_URL}/api/events/{event1_id}/registrations?checked_in=true", headers=headers_a)
    assert res_filter_checked.status_code == 200
    assert all(r["checked_in"] is True for r in res_filter_checked.json())

    # 13. Registration stats endpoint verification
    res_stats = requests.get(f"{BASE_URL}/api/events/{event1_id}/registrations/stats", headers=headers_a)
    assert res_stats.status_code == 200
    stats_data = res_stats.json()
    assert stats_data["total"] >= 3
    assert stats_data["checked_in"] >= 1
    assert stats_data["cancelled"] >= 1
    print(f"[TEST 12 & 13 PASS] Search, filter, and registration statistics endpoints verified.")

    print("\n==================================================")
    print("ALL 13 ATTENDEE OPERATIONS HUB TESTS PASSED!")
    print("==================================================")

if __name__ == "__main__":
    run_tests()
